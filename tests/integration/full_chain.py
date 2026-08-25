"""Repeatable runtime integration test for the complete EvalForge workflow.

The integration Compose stack uses a deterministic LLM stub for M02-M04. The
evaluation success path defaults to the OpenAI-compatible configuration in the
repository-local, ignored ``.env`` file (the requested DeepSeek setup).
"""
from __future__ import annotations

import argparse
import json
import sys
import time
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DOCUMENT = ROOT / "data" / "template" / "documents" / "附件1：上海银行中小企业无还本续贷业务操作规程（2026年版）.docx"
QA_TEMPLATE = ROOT / "data" / "template" / "qa_pairs.json"


class IntegrationFailure(RuntimeError):
    """Raised when an integration assertion fails with endpoint context."""


def load_local_env() -> dict[str, str]:
    values: dict[str, str] = {}
    env_path = ROOT / ".env"
    if not env_path.exists():
        return values
    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def require_deepseek_config() -> dict[str, str]:
    env = load_local_env()
    api_base = env.get("LLM_API_BASE") or env.get("LLM_API_URL") or ""
    api_key = env.get("LLM_API_KEY") or ""
    model = env.get("LLM_MODEL") or ""
    placeholder = not api_key or api_key.startswith("sk-xxx") or "placeholder" in api_key.lower()
    if not api_base or not model or placeholder:
        raise IntegrationFailure(
            ".env 中的 DeepSeek 配置不可用：需要有效的 LLM_API_BASE、LLM_API_KEY、LLM_MODEL"
        )
    return {"api_base": api_base, "api_key": api_key, "model": model}


class Chain:
    def __init__(self, base_url: str, *, timeout: float = 180.0) -> None:
        self.client = httpx.Client(base_url=base_url.rstrip("/"), timeout=timeout)
        self.created: dict[str, int] = {}

    def close(self) -> None:
        self.client.close()

    @staticmethod
    def step(message: str) -> None:
        print(f"[integration] {message}")  # noqa: T201 - command-line test progress

    def request(self, method: str, path: str, **kwargs: Any) -> httpx.Response:
        response = self.client.request(method, path, **kwargs)
        if response.status_code >= 400:
            detail = response.text[:500].replace("\n", " ")
            raise IntegrationFailure(f"{method} {path} -> {response.status_code}: {detail}")
        return response

    def json(self, method: str, path: str, **kwargs: Any) -> Any:
        response = self.request(method, path, **kwargs)
        try:
            return response.json()
        except ValueError as exc:
            raise IntegrationFailure(f"{method} {path} 未返回 JSON") from exc

    def wait_health(self, *, timeout: float = 120.0) -> dict[str, Any]:
        deadline = time.monotonic() + timeout
        last_error: Exception | None = None
        while time.monotonic() < deadline:
            try:
                adapters = self.json("GET", "/api/adapters")
                if isinstance(adapters, list) and adapters:
                    return {"status": "ok", "adapters": len(adapters)}
            except (httpx.HTTPError, IntegrationFailure) as exc:
                last_error = exc
            time.sleep(2)
        raise IntegrationFailure(f"联调入口在 {timeout:.0f}s 内未就绪: {last_error}")

    def poll_job(self, job_id: int, *, timeout: float = 300.0) -> dict[str, Any]:
        deadline = time.monotonic() + timeout
        last_progress = -1
        while time.monotonic() < deadline:
            job = self.json("GET", f"/api/jobs/{job_id}")
            progress = int(job.get("progress") or 0)
            if progress < last_progress:
                raise IntegrationFailure(f"job {job_id} 进度回退: {last_progress} -> {progress}")
            last_progress = progress
            if job.get("status") in {"completed", "done"}:
                return job
            if job.get("status") == "failed":
                raise IntegrationFailure(f"job {job_id} 失败: {job.get('message')}")
            time.sleep(1)
        raise IntegrationFailure(f"job {job_id} 在 {timeout:.0f}s 内未完成")

    def poll_run(self, run_id: int, *, timeout: float = 900.0) -> dict[str, Any]:
        deadline = time.monotonic() + timeout
        last_progress = -1
        while time.monotonic() < deadline:
            payload = self.json("GET", f"/api/evaluation-runs/{run_id}")
            run = payload.get("run") or {}
            progress = int(run.get("progress") or 0)
            if progress < last_progress:
                raise IntegrationFailure(f"run {run_id} 进度回退: {last_progress} -> {progress}")
            last_progress = progress
            if run.get("status") == "done":
                if progress != 100:
                    raise IntegrationFailure(f"run {run_id} 已完成但进度不是 100%")
                return payload
            if run.get("status") == "failed":
                raise IntegrationFailure(f"run {run_id} 失败")
            time.sleep(1)
        raise IntegrationFailure(f"run {run_id} 在 {timeout:.0f}s 内未完成")

    def upload_and_extract(self, document: Path) -> int:
        self.step(f"上传并解析 data 文档：{document.name}")
        content = document.read_bytes()
        files = {"file": (document.name, content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        precheck = self.json("POST", "/api/documents/precheck", files=files)
        if precheck.get("status") not in {"ok", "duplicate"}:
            raise IntegrationFailure(f"上传预检状态异常: {precheck}")
        if precheck.get("status") == "duplicate":
            raise IntegrationFailure("独立联调库不是空库：测试文档已存在，请重建联调卷")
        upload = self.json(
            "POST",
            "/api/documents/upload",
            files=files,
            data={"upload_user": "integration", "purpose": "basic"},
        )
        document_id = int(upload["document_id"])
        documents = self.json("GET", "/api/documents")
        stored = next((item for item in documents if int(item.get("document_id") or 0) == document_id), None)
        if not stored or stored.get("file_name") != document.name:
            raise IntegrationFailure("中文文档名未原样返回")
        extract = self.json("POST", "/api/eiu/extract", params={"document_id": document_id})
        if extract.get("status") not in {"completed", "done"}:
            self.poll_job(int(extract["job_id"]))
        eius = self.json("GET", f"/api/eiu/document/{document_id}")
        items = eius.get("items") if isinstance(eius, dict) else eius
        if not items or not any(item.get("is_questionable") for item in items):
            raise IntegrationFailure("知识点抽取完成，但没有可出题 EIU")
        self.created["document_id"] = document_id
        return document_id

    def generate_and_freeze(self, document_id: int) -> int:
        self.step("生成问答、执行质量门禁并冻结生成库版本")
        generated = self.json(
            "POST",
            "/api/cases/generate",
            params={"document_id": document_id},
            json={"angles": ["primary"], "include_variations": False, "dry_run": False},
        )
        if int(generated.get("generated") or 0) < 1:
            raise IntegrationFailure(f"问答生成数量异常: {generated}")
        quality = self.json("POST", "/api/quality-check", params={"document_id": document_id})
        if quality.get("failed_cases") or quality.get("errors"):
            raise IntegrationFailure(f"质量门禁存在失败题: {quality}")
        version = self.json(
            "POST",
            "/api/freeze",
            json={
                "name": "IT-生成库版本",
                "created_by": "integration",
                "document_ids": [document_id],
                "generation_config": {"source": "runtime-integration"},
            },
        )
        if version.get("status") != "frozen" or int(version.get("case_count") or 0) < 1:
            raise IntegrationFailure(f"冻结版本异常: {version}")
        version_id = int(version["version_id"])
        self.created["version_id"] = version_id
        return version_id

    @staticmethod
    def template_cases(count: int = 2) -> list[dict[str, Any]]:
        source = json.loads(QA_TEMPLATE.read_text(encoding="utf-8"))
        cases: list[dict[str, Any]] = []
        for item in source[:count]:
            cases.append(
                {
                    "q": item["question"],
                    "a": item["gold_answer"],
                    "dimension": item.get("question_type") or "general",
                    "evidence": item.get("evidence") or [],
                }
            )
        return cases

    def create_sources_and_composition(self, version_id: int) -> int:
        self.step("从 data 模板建立上传库、公共库，并合并三类来源")
        cases = self.template_cases()
        uploaded = self.json(
            "POST",
            "/api/eval-sets/upload",
            json={
                "name": "IT-上传库",
                "template_type": "single",
                "dimension": "rule",
                "source_file": "data/template/qa_pairs.json",
                "folder_path": "联调",
                "cases": cases,
            },
        )
        public = self.json(
            "POST",
            "/api/public-sets",
            json={"name": "IT-公共库", "version": "it-v1", "dimensions": ["rule"], "cases": cases},
        )
        uploaded_id = int(uploaded["set_id"])
        public_id = int(public["set_id"])
        composition = self.json(
            "POST",
            "/api/compositions",
            json={
                "name": "IT-三来源可执行评测集",
                "created_by": "integration",
                "items": [
                    {"source": "doc_generated", "version_id": version_id},
                    {"source": "uploaded", "set_id": uploaded_id},
                    {"source": "public", "set_id": public_id},
                ],
            },
        )
        composition_id = int(composition["composition_id"])
        resolved = self.json("GET", f"/api/compositions/{composition_id}")
        sources = {sample.get("source") for sample in resolved.get("samples") or []}
        if sources != {"doc_generated", "uploaded", "public"}:
            raise IntegrationFailure(f"组合未包含三类来源: {sorted(sources)}")
        self.created.update(
            {"uploaded_set_id": uploaded_id, "public_set_id": public_id, "composition_id": composition_id}
        )
        return composition_id

    def run_evaluation(self, composition_id: int, adapter: str) -> int:
        self.step("发起批量评测并轮询进度")
        if adapter == "deepseek":
            adapter_name = "openai_compatible"
            adapter_config = require_deepseek_config()
            run_name = "IT-DeepSeek-真实评测"
        else:
            adapter_name = "http"
            adapter_config = {
                "url": "http://integration-stub:8090/agent",
                "method": "POST",
                "headers": {"Content-Type": "application/json"},
                "body_template": '{"question":"{{question}}"}',
                "answer_path": "answer",
                "timeout_seconds": 10,
            }
            run_name = "IT-确定性联调评测"
        created = self.json(
            "POST",
            "/api/evaluation-runs",
            json={
                "composition_id": composition_id,
                "name": run_name,
                "adapter": adapter_name,
                "adapter_config": adapter_config,
            },
        )
        run_id = int(created["run_id"])
        self.poll_run(run_id)
        results = self.json("GET", f"/api/evaluation-runs/{run_id}/results")
        rows = results.get("results") or []
        summary = results.get("summary") or {}
        if len(rows) != int(created.get("total") or 0) or summary.get("total") != len(rows):
            raise IntegrationFailure("评测结果数量与运行样本总数不一致")
        if int(summary.get("error_count") or 0) != 0:
            raise IntegrationFailure(f"评测运行产生调用错误: error_count={summary.get('error_count')}")
        self.created["run_id"] = run_id
        return run_id

    def verify_export(self, run_id: int) -> None:
        self.step("验证完整与筛选后的 Excel 原始报告")
        results = self.json("GET", f"/api/evaluation-runs/{run_id}/results").get("results") or []
        full = self.request("POST", f"/api/evaluation-runs/{run_id}/export", json={"result_ids": None})
        workbook = load_workbook(BytesIO(full.content), read_only=True)
        if len(workbook.sheetnames) < 2 or workbook[workbook.sheetnames[0]].max_row != len(results) + 1:
            raise IntegrationFailure("完整 Excel 报告工作表或行数异常")
        selected_ids = [row["result_id"] for row in results[:1]]
        filtered = self.request(
            "POST", f"/api/evaluation-runs/{run_id}/export", json={"result_ids": selected_ids}
        )
        filtered_book = load_workbook(BytesIO(filtered.content), read_only=True)
        if filtered_book[filtered_book.sheetnames[0]].max_row != 2:
            raise IntegrationFailure("筛选 Excel 报告未严格限制为选中行")

    def verify_agent_failure(self, composition_id: int) -> None:
        self.step("验证目标智能体错误 JSON 不会卡死整轮评测")
        created = self.json(
            "POST",
            "/api/evaluation-runs",
            json={
                "composition_id": composition_id,
                "name": "IT-异常响应",
                "adapter": "http",
                "adapter_config": {
                    "url": "http://integration-stub:8090/agent?mode=invalid-json",
                    "method": "POST",
                    "body_template": '{"question":"{{question}}"}',
                    "answer_path": "answer",
                    "timeout_seconds": 2,
                },
            },
        )
        run_id = int(created["run_id"])
        self.poll_run(run_id)
        summary = self.json("GET", f"/api/evaluation-runs/{run_id}/results").get("summary") or {}
        if int(summary.get("error_count") or 0) < 1:
            raise IntegrationFailure("无效 JSON 未被记录为评测错误")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run the complete EvalForge integration chain")
    parser.add_argument("--base-url", default="http://127.0.0.1:18080")
    parser.add_argument("--document", type=Path, default=DEFAULT_DOCUMENT)
    parser.add_argument(
        "--evaluation-adapter",
        choices=("deepseek", "stub"),
        default="deepseek",
        help="deepseek reads the ignored .env; stub is for deterministic offline diagnosis",
    )
    parser.add_argument("--skip-failure-case", action="store_true")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.document.is_file():
        raise IntegrationFailure(f"测试文档不存在: {args.document}")
    chain = Chain(args.base_url)
    try:
        chain.wait_health()
        document_id = chain.upload_and_extract(args.document)
        version_id = chain.generate_and_freeze(document_id)
        composition_id = chain.create_sources_and_composition(version_id)
        run_id = chain.run_evaluation(composition_id, args.evaluation_adapter)
        chain.verify_export(run_id)
        if not args.skip_failure_case:
            chain.verify_agent_failure(composition_id)
        chain.step("PASS：完整主链路、三来源组合、评测结果与报告导出均已调通")
        return 0
    finally:
        chain.close()


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except IntegrationFailure as exc:
        print(f"[integration] FAIL: {exc}", file=sys.stderr)  # noqa: T201 - CLI failure output
        raise SystemExit(1) from exc
