import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook

from modules.m08_auto_evaluation import api
from modules.m08_auto_evaluation.api import _build_evaluation_workbook
from modules.m08_auto_evaluation.schemas import EvaluationRunRequest
from modules.m08_auto_evaluation.services.intermediate_metrics import aggregate_intermediate
from modules.m08_auto_evaluation.services.metrics import aggregate, score_case
from modules.m08_auto_evaluation.services.method_config import normalize_evaluation_selection
from modules.m08_auto_evaluation.services.text_metrics import MAX_ROUGE_L_CELLS, bleu_score, rouge_l_score, score_reference_metric
from modules.shared.services.database import DatabaseService


class M08MethodSelectionTests(unittest.TestCase):
    def test_legacy_method_selection_preserves_answer_rules_and_enabled_judge(self):
        self.assertEqual(
            normalize_evaluation_selection(None, None, legacy_judge_enabled=True),
            {
                "task_profile": "question_answering",
                "evaluation_methods": ["answer_comparison", "rules", "llm_as_judge"],
            },
        )

    def test_task_specific_methods_reject_mismatched_task_profile(self):
        with self.assertRaisesRegex(ValueError, "BLEU 仅适用于翻译"):
            normalize_evaluation_selection("question_answering", ["bleu"])
        with self.assertRaisesRegex(ValueError, "ROUGE-L 仅适用于文本生成"):
            normalize_evaluation_selection("translation", ["rouge_l"])

    def test_run_schema_exposes_task_methods_nodes_and_judge_options(self):
        request = EvaluationRunRequest(
            composition_id=3,
            task_profile="translation",
            evaluation_methods=["bleu", "llm_as_judge"],
            intermediate_eval={"enabled": True, "nodes": ["rag"]},
            judge_eval={"enabled": True, "prompt": "按标准答案判断", "include_retrieved": True},
        )
        self.assertEqual(request.task_profile, "translation")
        self.assertEqual(request.evaluation_methods, ["bleu", "llm_as_judge"])
        self.assertEqual(request.intermediate_eval.nodes, ["rag"])
        self.assertTrue(request.judge_eval.include_retrieved)

    def test_historical_run_serialization_defaults_and_masks_credentials(self):
        old_run = SimpleNamespace(
            run_id=8, composition_id=2, name="legacy", adapter="http",
            adapter_config={"api_key": "secret"}, intermediate_eval=None,
            judge_eval={"enabled": True, "prompt": "judge", "include_history": False,
                        "include_intermediate": False, "include_retrieved": False},
            task_profile=None, evaluation_methods=None, status="done", progress=100,
            total=1, finished=1, started_at=None, finished_at=None, created_at=None,
        )
        rendered = DatabaseService._evaluation_run_to_dict(old_run)
        self.assertEqual(rendered["task_profile"], "question_answering")
        self.assertEqual(rendered["evaluation_methods"], ["answer_comparison", "rules", "llm_as_judge"])
        self.assertEqual(rendered["adapter_config"]["api_key"], "***")

    def test_run_api_persists_and_dispatches_selected_methods(self):
        request = EvaluationRunRequest(
            composition_id=3,
            task_profile="translation",
            evaluation_methods=["bleu"],
        )
        with (
            patch.object(api._db, "get_composition", return_value={"composition_id": 3}),
            patch.object(api, "resolve_composition", return_value=[{"case_uid": "case-1"}]),
            patch.object(api, "get_adapter", return_value=object()),
            patch.object(api._db, "save_evaluation_run", return_value=21) as save_run,
            patch.object(api._db, "save_audit"),
            patch.object(api, "start_run_async") as start_run,
        ):
            response = api.create_evaluation_run(request)
        self.assertEqual(response["run_id"], 21)
        self.assertEqual(save_run.call_args.kwargs["task_profile"], "translation")
        self.assertEqual(save_run.call_args.kwargs["evaluation_methods"], ["bleu"])
        self.assertEqual(start_run.call_args.kwargs["evaluation_methods"], ["bleu"])

    def test_run_api_rejects_incompatible_task_metric_before_persisting(self):
        request = EvaluationRunRequest(
            composition_id=3,
            task_profile="question_answering",
            evaluation_methods=["bleu"],
        )
        with (
            patch.object(api._db, "get_composition", return_value={"composition_id": 3}),
            patch.object(api, "resolve_composition", return_value=[{"case_uid": "case-1"}]),
            patch.object(api, "get_adapter", return_value=object()),
            patch.object(api._db, "save_evaluation_run") as save_run,
        ):
            with self.assertRaises(HTTPException) as caught:
                api.create_evaluation_run(request)
        self.assertEqual(caught.exception.status_code, 400)
        save_run.assert_not_called()

    def test_reference_metrics_compute_and_missing_data_is_explicit(self):
        self.assertEqual(bleu_score("hello world", ["hello world"]), 1.0)
        self.assertEqual(rouge_l_score("a b c", ["a b c"])["score"], 1.0)
        scores = score_case(
            {"gold_answer": "参考文本"},
            {"answer": "模型回答", "usage": {}},
            evaluation_methods=["bleu"],
            task_profile="translation",
        )
        self.assertEqual(scores["bleu"]["status"], "scored")
        self.assertIsNone(scores["score"])
        missing = score_case(
            {"question": "问题"},
            {"answer": "回答", "usage": {}},
            evaluation_methods=["rouge_l"],
            task_profile="text_generation",
        )
        self.assertEqual(missing["rouge_l"]["status"], "data_missing")
        self.assertIsNone(missing["rouge_l"]["score"])

    def test_rouge_l_marks_extreme_inputs_unavailable_instead_of_blocking(self):
        tokens_per_side = int(MAX_ROUGE_L_CELLS ** 0.5) + 1
        long_text = " ".join(["x"] * tokens_per_side)
        result = score_reference_metric("rouge_l", {"gold_answer": long_text}, long_text)
        self.assertEqual(result["status"], "unavailable")
        self.assertIsNone(result["score"])
        self.assertIn("计算上限", result["reason"])

    def test_selected_methods_run_independently_and_rule_missing_is_not_zero(self):
        scores = score_case(
            {"gold_answer": "正确答案"},
            {"answer": "正确答案", "usage": {}},
            evaluation_methods=["rules"],
        )
        self.assertIsNone(scores["score"])
        self.assertEqual(scores["rule"]["status"], "data_missing")
        self.assertIsNone(scores["rule"]["score"])

    def test_selected_method_errors_remain_distinguishable_from_not_selected(self):
        scores = score_case(
            {"question": "q", "gold_answer": "ref"},
            {"error": "timeout", "usage": {}},
            judge_config={"enabled": True, "prompt": "judge"},
            evaluation_methods=["answer_comparison", "rules", "llm_as_judge", "bleu"],
            task_profile="translation",
        )
        self.assertEqual(scores["answer_comparison"]["status"], "error")
        self.assertEqual(scores["rule"]["status"], "error")
        self.assertEqual(scores["judge"]["status"], "error")
        self.assertEqual(scores["bleu"]["status"], "error")
        self.assertIn("max_order", scores["bleu"]["parameters"])

    def test_intent_aggregate_adds_standard_micro_metrics_and_confusion(self):
        results = [
            {"scores": {"intermediate": {"nodes": {"intent": {"status": "scored", "reference": "a", "predicted": "a", "correct": True}}}}},
            {"scores": {"intermediate": {"nodes": {"intent": {"status": "scored", "reference": "b", "predicted": "a", "correct": False}}}}},
        ]
        intent = aggregate_intermediate(results, ["intent"])["nodes"]["intent"]
        self.assertEqual(intent["accuracy"], 0.5)
        self.assertEqual(intent["micro_precision"], 0.5)
        self.assertEqual(intent["micro_recall"], 0.5)
        self.assertEqual(intent["micro_f1"], 0.5)
        self.assertEqual(intent["macro_f1"], 0.3333)
        self.assertEqual(intent["confusion"], {"a": {"a": 1}, "b": {"a": 1}})

    def test_unavailable_intermediate_dependency_is_not_reported_as_missing_or_zero(self):
        summary = aggregate_intermediate(
            [{"scores": {"intermediate": {"nodes": {
                "rag": {"status": "unavailable", "metrics": {"context_recall": None}}
            }}}}],
            ["rag"],
        )
        self.assertEqual(summary["nodes"]["rag"]["status"], "unavailable")
        self.assertIsNone(summary["nodes"]["rag"]["context_recall"])

    def test_run_summary_contains_task_selection_and_reference_metric(self):
        scores = score_case(
            {"gold_answer": "the cat sat"},
            {"answer": "the cat sat", "usage": {}},
            evaluation_methods=["bleu"],
            task_profile="translation",
        )
        summary = aggregate(
            [{"scores": scores}],
            task_profile="translation",
            evaluation_methods=["bleu"],
        )
        self.assertEqual(summary["task_profile"], "translation")
        self.assertEqual(summary["evaluation_methods"], ["bleu"])
        self.assertEqual(summary["reference_metrics"]["bleu"]["score"], 1.0)
        self.assertEqual(summary["reference_metrics"]["bleu"]["scored"], 1)
        partial = aggregate(
            [
                {"scores": {"rouge_l": {"status": "scored", "score": 0.5}}},
                {"scores": {"rouge_l": {"status": "unavailable", "score": None}}},
            ],
            task_profile="text_generation",
            evaluation_methods=["rouge_l"],
        )
        self.assertEqual(partial["reference_metrics"]["rouge_l"]["unavailable"], 1)

    def test_workbook_exports_run_parameters_national_labels_and_node_metrics(self):
        content = _build_evaluation_workbook(
            {
                "run_id": 4,
                "composition_id": 9,
                "task_profile": "translation",
                "evaluation_methods": ["bleu"],
                "intermediate_eval": {"enabled": True, "nodes": ["intent"]},
                "judge_eval": {"enabled": False},
            },
            [{
                "question": "Translate this",
                "gold_answer": "你好",
                "answer": "你好",
                "scores": {
                    "score": None,
                    "bleu": {"status": "scored", "score": 1.0, "parameters": {"max_order": 4}},
                    "intermediate": {"nodes": {"intent": {
                        "status": "scored", "reference": "greeting", "predicted": "greeting", "correct": True,
                    }}},
                },
            }],
        )
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        self.assertIn("指标汇总", workbook.sheetnames)
        self.assertIn("意图分类明细", workbook.sheetnames)
        config_rows = list(workbook["运行摘要"].iter_rows(values_only=True))
        self.assertIn(("任务类型", "翻译"), config_rows)
        metric_rows = list(workbook["指标汇总"].iter_rows(values_only=True))
        bleu_row = next(row for row in metric_rows if row[0] == "BLEU")
        self.assertEqual(bleu_row[2], 1.0)
        self.assertEqual(bleu_row[3], "已计算")
        self.assertIn("GB/T 45288.2—2025", bleu_row[6])
        raw_headers = next(workbook["原始评测报告"].iter_rows(min_row=1, max_row=1, values_only=True))
        self.assertIn("意图标准标签", raw_headers)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
