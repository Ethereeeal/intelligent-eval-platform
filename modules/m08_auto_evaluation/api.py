"""M08 — Agent 评测 API（BRD §9）。

路由：
  POST /api/evaluation-runs                  发起批量运行（202，异步）
  GET  /api/evaluation-runs                  运行列表
  GET  /api/evaluation-runs/{id}             运行进度 + 汇总
  GET  /api/evaluation-runs/{id}/results     单题结果 + 分层指标汇总
  GET  /api/evaluation-runs/{id}/failures    E2E/D9 失败记录（ErrorBook）
  POST /api/evaluation-runs/{id}/cancel      取消运行（当前单题结束后生效）
  POST /api/evaluation-results/{id}/retry    单题复测（保留尝试链）
  PATCH /api/error-book/{id}                 人工处置异常项
  GET  /api/error-book                       智能体失败诊断与优化分析数据源
  GET  /api/adapters                         内置适配器清单
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime
import json
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.m05_dataset_lifecycle.services.composition import resolve_composition
from modules.m08_auto_evaluation.schemas import (
    AdapterTestRequest,
    ContextAnalysisRequest,
    ErrorBookUpdateRequest,
    EvaluationCaseRetryRequest,
    EvaluationExportRequest,
    EvaluationRunRequest,
)
from modules.m08_auto_evaluation.services.adapter import (
    ADAPTER_REGISTRY,
    AdapterError,
    get_adapter,
)
from modules.m08_auto_evaluation.services.intermediate_metrics import normalize_intermediate_config
from modules.m08_auto_evaluation.services.judge import normalize_judge_config
from modules.m08_auto_evaluation.services.method_config import (
    METHOD_LABELS,
    METHOD_STANDARD_LABELS,
    TASK_PROFILE_LABELS,
    normalize_evaluation_selection,
)
from modules.m08_auto_evaluation.services.metrics import aggregate
from modules.m08_auto_evaluation.services.optimization import cluster_error_book
from modules.m08_auto_evaluation.services.runner import start_case_retry_async, start_run_async
from modules.shared.services.database import DatabaseService

evaluation_router = APIRouter(prefix="/api", tags=["m08-auto-evaluation"])

_db = DatabaseService()

_EVALUATION_STATUS_LABELS = {
    "scored": "已计算",
    "data_missing": "不适用：缺少数据",
    "unavailable": "不可用：依赖或模型未配置",
    "error": "计算失败",
    "not_selected": "未选择",
}


def _evaluation_status_label(value: str | None) -> str:
    return _EVALUATION_STATUS_LABELS.get(value or "", value or "未计算")


def _sanitize_adapter_config(config: dict | None) -> dict | None:
    """持久化前剔除敏感字段（API Key 不进库，避免经接口回显泄露）。"""
    if not config:
        return config
    sanitized = dict(config)
    sanitized.pop("api_key", None)
    # 通用 HTTP 模式的 Header 常携带 Token；运行记录只保留字段名，绝不落库其值。
    if isinstance(sanitized.get("headers"), dict):
        sanitized["headers"] = {key: "***" for key in sanitized["headers"]}
    return sanitized


@evaluation_router.post("/evaluation-runs", status_code=202)
def create_evaluation_run(payload: EvaluationRunRequest):
    """发起批量运行：组合解析为统一输入样本 → 异步线程逐题调用适配器。"""
    if _db.get_composition(payload.composition_id) is None:
        raise HTTPException(status_code=404, detail="composition not found")
    try:
        samples = resolve_composition(_db, payload.composition_id)
        adapter = get_adapter(payload.adapter, payload.adapter_config)
        intermediate_config = normalize_intermediate_config(payload.intermediate_eval.model_dump())
        selection = normalize_evaluation_selection(
            payload.task_profile,
            payload.evaluation_methods,
            legacy_judge_enabled=payload.judge_eval.enabled,
        )
        judge_payload = payload.judge_eval.model_dump()
        judge_payload["enabled"] = "llm_as_judge" in selection["evaluation_methods"]
        judge_config = normalize_judge_config(judge_payload)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except AdapterError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not samples:
        raise HTTPException(status_code=400, detail="组合解析为空，无可评测样本")
    if not selection["evaluation_methods"] and not intermediate_config["enabled"]:
        raise HTTPException(status_code=400, detail="至少选择一种评测方法或一个中间评测节点")
    run_id = _db.save_evaluation_run(
        composition_id=payload.composition_id,
        name=payload.name,
        adapter=payload.adapter,
        adapter_config=_sanitize_adapter_config(payload.adapter_config),
        intermediate_eval=intermediate_config,
        judge_eval=judge_config,
        task_profile=selection["task_profile"],
        evaluation_methods=selection["evaluation_methods"],
    )
    start_run_async(
        run_id=run_id,
        samples=samples,
        adapter=adapter,
        intermediate_config=intermediate_config,
        judge_config=judge_config,
        evaluation_methods=selection["evaluation_methods"],
        task_profile=selection["task_profile"],
    )
    try:
        _db.save_audit(
            operation="evaluation_run.create",
            target_type="evaluation_run",
            target_id=str(run_id),
            actor="web",
            detail={
                "composition_id": payload.composition_id,
                "adapter": payload.adapter,
                "total": len(samples),
                "intermediate_eval": intermediate_config,
                "judge_eval": judge_config,
                "task_profile": selection["task_profile"],
                "evaluation_methods": selection["evaluation_methods"],
            },
        )
    except Exception:  # noqa: BLE001 — 审计失败不阻断
        pass
    return {"run_id": run_id, "status": "running", "total": len(samples)}


@evaluation_router.get("/evaluation-runs")
def list_evaluation_runs():
    return _db.list_evaluation_runs()


@evaluation_router.delete("/evaluation-runs/{run_id}")
def delete_evaluation_run(run_id: int, confirm: bool = Query(default=False)):
    """永久删除终态运行；必须显式二次确认，运行中任务应先取消。"""
    run = _get_run_or_404(run_id)
    if run.get("status") in {"pending", "running", "cancelling"}:
        raise HTTPException(status_code=409, detail="运行中任务不能删除，请先取消")
    if not confirm:
        raise HTTPException(status_code=400, detail="永久删除需要二次确认")
    deleted = _db.delete_evaluation_run(run_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="evaluation run not found")
    try:
        _db.save_audit(
            operation="evaluation_run.delete",
            target_type="evaluation_run",
            target_id=str(run_id),
            actor="web",
        )
    except Exception:  # noqa: BLE001 — 审计失败不阻断
        pass
    return {"ok": True, "run_id": run_id}


@evaluation_router.post("/evaluation-runs/{run_id}/cancel", status_code=202)
def cancel_evaluation_run(run_id: int):
    run = _get_run_or_404(run_id)
    if run.get("status") not in {"pending", "running", "cancelling"}:
        raise HTTPException(status_code=409, detail="当前运行已结束，不能取消")
    status = "cancelled" if run.get("status") == "pending" else "cancelling"
    updated = _db.update_evaluation_run(run_id, status=status)
    return {"run_id": run_id, "status": updated["status"]}


def _get_run_or_404(run_id: int) -> dict:
    run = _db.get_evaluation_run(run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="evaluation run not found")
    return run


@evaluation_router.get("/evaluation-runs/{run_id}")
def get_evaluation_run(run_id: int):
    run = _get_run_or_404(run_id)
    results = _db.list_evaluation_results(run_id)
    return {
        "run": run,
        "summary": aggregate(
            results,
            run.get("intermediate_eval"),
            run.get("judge_eval"),
            run.get("task_profile"),
            run.get("evaluation_methods"),
        ),
    }


@evaluation_router.get("/evaluation-runs/{run_id}/results")
def evaluation_run_results(run_id: int):
    run = _get_run_or_404(run_id)
    results = _db.list_evaluation_results(run_id)
    issues = _db.list_error_book(run_id=run_id)
    by_result = {item.get("result_id"): item for item in issues if item.get("result_id")}
    by_case = {}
    for issue in issues:  # list_error_book 按新到旧排序，保留每道题最新处置记录
        if issue.get("case_uid"):
            by_case.setdefault(issue["case_uid"], issue)
    for item in results:
        item["error_book"] = by_result.get(item["result_id"]) or by_case.get(item.get("case_uid"))
    return {
        "results": results,
        "summary": aggregate(
            results,
            run.get("intermediate_eval"),
            run.get("judge_eval"),
            run.get("task_profile"),
            run.get("evaluation_methods"),
        ),
    }


@evaluation_router.get("/evaluation-results/{result_id}/attempts")
def evaluation_result_attempts(result_id: int):
    if _db.get_evaluation_result(result_id) is None:
        raise HTTPException(status_code=404, detail="evaluation result not found")
    return {"attempts": _db.list_evaluation_result_attempts(result_id)}


@evaluation_router.patch("/evaluation-results/{result_id}/context-analysis")
def update_context_analysis(result_id: int, payload: ContextAnalysisRequest):
    """保存失败/不确定多轮样本的上下文分析，不覆盖原始评分。"""
    result = _db.get_evaluation_result(result_id)
    if result is None:
        raise HTTPException(status_code=404, detail="evaluation result not found")
    if not result.get("turn_trace"):
        raise HTTPException(status_code=400, detail="该结果没有可分析的多轮请求上下文")
    if not payload.note.strip():
        raise HTTPException(status_code=400, detail="上下文分析备注不能为空")
    if result.get("status") == "passed" and (result.get("scores") or {}).get("score") is not None:
        raise HTTPException(status_code=400, detail="通过样本无需提交上下文错误分析")
    analysis = {
        "category": payload.category,
        "note": payload.note.strip(),
        "updated_at": datetime.utcnow().isoformat(),
    }
    updated = _db.update_evaluation_case_result(result_id, context_analysis=analysis)
    issue = _db.get_error_book_item_for_result(
        result_id,
        run_id=result.get("run_id"),
        case_uid=result.get("case_uid"),
    )
    if issue:
        _db.update_error_book_item(
            issue["item_id"],
            root_cause=f"{payload.category}: {payload.note.strip()}",
        )
    return updated


@evaluation_router.post("/evaluation-results/{result_id}/retry", status_code=202)
def retry_evaluation_result(result_id: int, payload: EvaluationCaseRetryRequest):
    result = _db.get_evaluation_result(result_id)
    if result is None:
        raise HTTPException(status_code=404, detail="evaluation result not found")
    if result.get("parent_result_id") is not None:
        raise HTTPException(status_code=400, detail="请从原始结果发起复测")
    run = _get_run_or_404(result["run_id"])
    config = dict(run.get("adapter_config") or {})
    config.update(payload.adapter_config or {})
    try:
        adapter = get_adapter(run["adapter"], config)
    except (ValueError, AdapterError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    attempt_id = start_case_retry_async(
        run_id=run["run_id"],
        source_result=result,
        adapter=adapter,
        analysis_threshold=payload.analysis_threshold,
        intermediate_config=run.get("intermediate_eval"),
        judge_config=run.get("judge_eval"),
        evaluation_methods=run.get("evaluation_methods"),
        task_profile=run.get("task_profile") or "question_answering",
    )
    return {"result_id": result_id, "attempt_result_id": attempt_id, "status": "pending"}


def _build_evaluation_workbook(run: dict, results: list[dict]) -> bytes:
    """生成便于测试人员筛选和归因的 Excel 原始报告。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "原始评测报告"
    headers = [
        "序号", "问题", "标准答案", "智能体回答", "得分", "规则评测", "LLM Judge评分", "LLM Judge理由", "状态", "耗时(ms)",
        "维度", "难度", "来源", "归因", "错误信息", "用例ID", "智能体额外输出",
        "标准答案比对状态", "标准答案比对方法", "规则评测状态", "规则评测明细(JSON)", "Judge状态", "Judge判据(JSON)",
        "BLEU", "BLEU状态", "ROUGE-L", "ROUGE-L Precision", "ROUGE-L Recall", "ROUGE-L状态", "改写语义相似度", "改写约束Precision", "改写约束Recall", "改写约束F1",
        "意图标准标签", "意图预测标签", "意图是否正确", "RAG Context Precision", "RAG Context Recall", "RAG Faithfulness", "RAG Answer Relevancy",
        "BLEU计算参数(JSON)", "ROUGE-L计算参数(JSON)", "中间节点输出(JSON)",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="6750A4")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, item in enumerate(results, start=1):
        scores = item.get("scores") or {}
        intermediate_nodes = ((scores.get("intermediate") or {}).get("nodes") or {})
        rewrite = intermediate_nodes.get("rewrite") or {}
        intent = intermediate_nodes.get("intent") or {}
        rag = intermediate_nodes.get("rag") or {}
        rag_metrics = rag.get("metrics") or {}
        bleu = scores.get("bleu") or {}
        rouge_l = scores.get("rouge_l") or {}
        rule = scores.get("rule") or {}
        judge = scores.get("judge") or {}
        answer_comparison = scores.get("answer_comparison") or {}
        sheet.append([
            index,
            item.get("question") or "",
            item.get("gold_answer") or "",
            item.get("answer") or "",
            scores.get("score"),
            rule.get("score"),
            judge.get("score"),
            judge.get("reason"),
            item.get("status") or "",
            scores.get("latency_ms"),
            item.get("dimension") or "",
            item.get("difficulty") or "",
            item.get("source") or "",
            item.get("diagnosis") or "",
            item.get("error_message") or "",
            item.get("case_uid") or "",
            json.dumps(item.get("agent_observations"), ensure_ascii=False) if item.get("agent_observations") else "",
            _evaluation_status_label(answer_comparison.get("status", "scored" if scores.get("score") is not None else "not_selected")),
            answer_comparison.get("method", scores.get("method")),
            _evaluation_status_label(rule.get("status", "scored" if rule.get("score") is not None else "not_selected")),
            json.dumps(rule, ensure_ascii=False) if rule else "",
            _evaluation_status_label(judge.get("status", "scored" if judge.get("score") is not None else "not_selected")),
            json.dumps(judge.get("criteria") or [], ensure_ascii=False) if judge else "",
            bleu.get("score"), _evaluation_status_label(bleu.get("status", "not_selected")),
            rouge_l.get("score"), rouge_l.get("precision"), rouge_l.get("recall"), _evaluation_status_label(rouge_l.get("status", "not_selected")),
            rewrite.get("semantic_similarity"), rewrite.get("constraint_precision"),
            rewrite.get("constraint_recall"), rewrite.get("constraint_f1"),
            intent.get("reference"), intent.get("predicted"), intent.get("correct"),
            rag_metrics.get("context_precision"), rag_metrics.get("context_recall"),
            rag_metrics.get("faithfulness"), rag_metrics.get("answer_relevancy"),
            json.dumps(bleu.get("parameters") or {}, ensure_ascii=False) if bleu else "",
            json.dumps(rouge_l.get("parameters") or {}, ensure_ascii=False) if rouge_l else "",
            json.dumps(intermediate_nodes, ensure_ascii=False) if intermediate_nodes else "",
        ])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [8, 42, 42, 42, 10, 12, 14, 36, 12, 12, 16, 16, 16, 24, 32, 24, 40] + [18] * (len(headers) - 17)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24

    summary = workbook.create_sheet("运行摘要")
    composition_id = run.get("composition_id")
    task_profile = run.get("task_profile") or "question_answering"
    evaluation_methods = run.get("evaluation_methods")
    if evaluation_methods is None:
        evaluation_methods = ["answer_comparison", "rules"]
        if (run.get("judge_eval") or {}).get("enabled"):
            evaluation_methods.append("llm_as_judge")
    intermediate_config = run.get("intermediate_eval") or {"enabled": False, "nodes": []}
    judge_config = run.get("judge_eval") or {}
    method_labels = [METHOD_LABELS.get(item, item) for item in evaluation_methods]
    node_labels = {"rewrite": "改写", "intent": "意图识别", "rag": "RAG"}
    selected_nodes = intermediate_config.get("nodes") or []
    summary_metrics = aggregate(
        results,
        intermediate_config,
        judge_config,
        task_profile,
        evaluation_methods,
    )
    summary_rows = [
        ("运行名称", run.get("name") or f"运行 #{run.get('run_id')}"),
        ("运行ID", run.get("run_id")),
        ("可执行评测集版本", f"#{composition_id}" if composition_id is not None else "—"),
        ("适配器", run.get("adapter") or ""),
        ("任务类型", TASK_PROFILE_LABELS.get(task_profile, task_profile)),
        ("评测方法", "、".join(method_labels) or "未选择"),
        ("中间过程节点", "、".join(node_labels.get(item, item) for item in selected_nodes) or "未启用"),
        ("Judge附加输入", "、".join(label for enabled, label in [
            (judge_config.get("include_history"), "多轮历史"),
            (judge_config.get("include_intermediate"), "中间节点结果"),
            (judge_config.get("include_retrieved"), "检索结果"),
        ] if enabled) or "无"),
        ("国标标注说明", "国标公式/条文引用不代表本项目通过国标认证；人工 MOS 本版本未纳入。"),
        ("状态", run.get("status") or ""),
        ("结果数量", len(results)),
        ("创建时间", run.get("created_at") or ""),
        ("开始时间", run.get("started_at") or ""),
        ("完成时间", run.get("finished_at") or ""),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 48
    for cell in summary[1]:
        cell.font = Font(bold=True)

    metric_sheet = workbook.create_sheet("指标汇总")
    metric_sheet.append(["节点 / 方法", "指标", "结果", "计算状态", "已评分样本", "适用样本", "国标标注", "计算参数 / 说明"])
    header_fill = PatternFill("solid", fgColor="6750A4")
    for cell in metric_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def add_metric(
        group: str,
        metric: str,
        value: object,
        scored: object,
        total: object,
        standard: str,
        params: object = "",
        status: str | None = None,
    ) -> None:
        metric_sheet.append([
            group, metric, value if value is not None else "不适用",
            _evaluation_status_label(status or ("scored" if value is not None else "data_missing")),
            scored, total, standard, params,
        ])
        if isinstance(value, (int, float)) and 0 <= value <= 1:
            metric_sheet.cell(row=metric_sheet.max_row, column=3).number_format = "0.0%"

    scored_answer = [item for item in results if (item.get("scores") or {}).get("score") is not None]
    add_metric(
        "标准答案比对", "平均得分",
        round(sum(float(item["scores"]["score"]) for item in scored_answer) / len(scored_answer), 4) if scored_answer else None,
        len(scored_answer), len(results), "项目方法", "短答案精确匹配；长答案语义相似度",
        "scored" if scored_answer else "data_missing",
    )
    if "rules" in evaluation_methods:
        add_metric("规则评测", "通过率", summary_metrics.get("rule_passed_rate"), summary_metrics.get("rule_scored"), len(results), "项目方法", "must_have_points / acceptable_answers", "scored" if summary_metrics.get("rule_scored") else "data_missing")
    if "llm_as_judge" in evaluation_methods:
        judge_statuses = Counter(
            ((item.get("scores") or {}).get("judge") or {}).get("status", "data_missing")
            for item in results
        )
        judge_status = (
            "scored" if summary_metrics.get("judge_scored") else "error" if judge_statuses.get("error")
            else "unavailable" if judge_statuses.get("unavailable") else "data_missing"
        )
        add_metric("LLM-as-a-Judge", "通过率", summary_metrics.get("judge_passed_rate"), summary_metrics.get("judge_scored"), len(results), "GB/T 45288.2—2025 第 6.5(c) 提及的方法（非指标）", "固定输入问题/标准答案/实际回答；可选上下文见运行摘要", judge_status)
    for method in ("bleu", "rouge_l"):
        if method in evaluation_methods:
            metric = summary_metrics.get("reference_metrics", {}).get(method) or {}
            params = next(
                (
                    ((item.get("scores") or {}).get(method) or {}).get("parameters")
                    for item in results
                    if ((item.get("scores") or {}).get(method) or {}).get("parameters")
                ),
                None,
            )
            add_metric(METHOD_LABELS[method], "平均分", metric.get("score"), metric.get("scored"), metric.get("total"), METHOD_STANDARD_LABELS[method], json.dumps(params or {}, ensure_ascii=False), metric.get("status"))

    intermediate = summary_metrics.get("intermediate") or {}
    for node, metrics in (intermediate.get("nodes") or {}).items():
        node_label = node_labels.get(node, node)
        if node == "rewrite":
            node_metrics = [
                ("语义相似度", "semantic_similarity", "项目方法"),
                ("约束Precision", "constraint_precision", "项目方法"),
                ("约束Recall", "constraint_recall", "项目方法"),
                ("约束F1", "constraint_f1", "项目方法"),
            ]
        elif node == "intent":
            node_metrics = [
                ("Accuracy", "accuracy", "GB/T 45288.2—2025 附录 A.1.1 公式"),
                ("Precision（micro）", "micro_precision", "GB/T 45288.2—2025 附录 A.1.3 公式"),
                ("Recall（micro）", "micro_recall", "GB/T 45288.2—2025 附录 A.1.2 公式"),
                ("Micro-F1", "micro_f1", "GB/T 45288.2—2025 附录 A.1.4 公式"),
                ("Macro-F1", "macro_f1", "项目扩展；非附录 A.1 明示指标"),
            ]
        else:
            node_metrics = [(name, name, "RAGAS 项目方法") for name in (
                "context_precision", "context_recall", "faithfulness", "answer_relevancy"
            )]
        for label, key, standard in node_metrics:
            value = metrics.get(key)
            metric_status = (
                "scored" if value is not None else metrics.get("status")
                if metrics.get("status") in {"error", "unavailable"} else "data_missing"
            )
            add_metric(node_label, label, value, metrics.get("scored"), len(results), standard, status=metric_status)
    metric_sheet.freeze_panes = "A2"
    metric_sheet.auto_filter.ref = metric_sheet.dimensions
    for index, width in enumerate([20, 24, 16, 18, 16, 16, 54, 72], start=1):
        metric_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in metric_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    intent_summary = (intermediate.get("nodes") or {}).get("intent") or {}
    if intent_summary.get("by_label"):
        intent_sheet = workbook.create_sheet("意图分类明细")
        intent_sheet.append(["标准意图标签", "Precision", "Recall", "F1（项目扩展）", "Support", "标准标注"])
        for cell in intent_sheet[1]:
            cell.fill = PatternFill("solid", fgColor="6750A4")
            cell.font = Font(color="FFFFFF", bold=True)
        for label, values in intent_summary["by_label"].items():
            intent_sheet.append([
                label, values.get("precision"), values.get("recall"), values.get("f1"), values.get("support"),
                "Precision / Recall 公式参照 GB/T 45288.2—2025 附录 A.1.3 / A.1.2；类别 F1 为项目扩展。",
            ])
        intent_sheet.append(["混淆矩阵", json.dumps(intent_summary.get("confusion") or {}, ensure_ascii=False), "", "", "", ""])
        intent_sheet.column_dimensions["A"].width = 28
        for column in "BCDE":
            intent_sheet.column_dimensions[column].width = 20
        intent_sheet.column_dimensions["F"].width = 78
        for row in intent_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    for exported_sheet in workbook.worksheets:
        for row in exported_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
    workbook.save(output)
    return output.getvalue()


@evaluation_router.post("/evaluation-runs/{run_id}/export")
def export_evaluation_run(run_id: int, payload: EvaluationExportRequest):
    run = _get_run_or_404(run_id)
    results = _db.list_evaluation_results(run_id)
    if payload.result_ids is not None:
        selected_ids = set(payload.result_ids)
        results = [item for item in results if item.get("result_id") in selected_ids]
    content = _build_evaluation_workbook(run, results)
    base_name = (run.get("name") or f"评测报告-{run_id}").replace("/", "-").replace("\\", "-")
    filename = f"{base_name}.xlsx"
    headers = {"Content-Disposition": f"attachment; filename=report-{run_id}.xlsx; filename*=UTF-8''{quote(filename)}"}
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@evaluation_router.get("/evaluation-runs/{run_id}/failures")
def evaluation_run_failures(run_id: int):
    _get_run_or_404(run_id)
    return _db.list_error_book(run_id=run_id)


@evaluation_router.post("/evaluation-runs/{run_id}/retry")
def retry_evaluation_run(run_id: int):
    """重跑（创建新 run，供同一冻结评测集复测比较，FR-OPT-002）。"""
    run = _get_run_or_404(run_id)
    if run.get("composition_id") is None:
        raise HTTPException(status_code=400, detail="原运行缺少组合，无法重跑")
    try:
        samples = resolve_composition(_db, run["composition_id"])
        adapter = get_adapter(run["adapter"], run.get("adapter_config"))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except AdapterError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not samples:
        raise HTTPException(status_code=400, detail="组合解析为空，无可评测样本")
    new_run_id = _db.save_evaluation_run(
        composition_id=run["composition_id"],
        name=(run.get("name") or "评测运行") + "（重跑）",
        adapter=run["adapter"],
        adapter_config=_sanitize_adapter_config(run.get("adapter_config")),
        intermediate_eval=run.get("intermediate_eval"),
        judge_eval=run.get("judge_eval"),
        task_profile=run.get("task_profile") or "question_answering",
        evaluation_methods=run.get("evaluation_methods"),
    )
    start_run_async(
        run_id=new_run_id,
        samples=samples,
        adapter=adapter,
        intermediate_config=run.get("intermediate_eval"),
        judge_config=run.get("judge_eval"),
        evaluation_methods=run.get("evaluation_methods"),
        task_profile=run.get("task_profile") or "question_answering",
    )
    return {"run_id": new_run_id, "status": "running", "total": len(samples)}


@evaluation_router.get("/error-book")
def error_book(
    diagnosis: str | None = Query(default=None, description="按 E2E/D9 过滤"),
    status: str | None = Query(default=None, description="open/processed/verified/ignored"),
):
    items = _db.list_error_book(diagnosis=diagnosis, status=status)
    return {"items": items, "clusters": cluster_error_book(items)}


@evaluation_router.patch("/error-book/{item_id}")
def update_error_book(item_id: int, payload: ErrorBookUpdateRequest):
    if payload.status == "ignored" and not (payload.resolution_note or "").strip():
        raise HTTPException(status_code=400, detail="忽略异常时必须填写原因")
    if payload.status == "processed" and not (payload.resolution_category or "").strip():
        raise HTTPException(status_code=400, detail="标记已处理时必须选择人工分类")
    item = _db.update_error_book_item(
        item_id,
        status=payload.status,
        resolution_category=payload.resolution_category,
        resolution_note=payload.resolution_note,
    )
    if item is None:
        raise HTTPException(status_code=404, detail="error book item not found")
    return item


@evaluation_router.post("/adapters/test")
def test_adapter(payload: AdapterTestRequest):
    try:
        adapter = get_adapter(payload.adapter, payload.adapter_config)
        result = adapter.run_single(payload.question)
    except (ValueError, AdapterError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if result.get("error"):
        raise HTTPException(status_code=502, detail=str(result["error"])[:300])
    return {
        "ok": True,
        "answer": result.get("answer"),
        "agent_observations": result.get("agent_observations"),
        "usage": result.get("usage") or {},
    }


@evaluation_router.get("/adapters")
def list_adapters():
    return [
        {
            "name": name,
            "description": (adapter_cls.__doc__ or "").strip(),
        }
        for name, adapter_cls in ADAPTER_REGISTRY.items()
    ]
