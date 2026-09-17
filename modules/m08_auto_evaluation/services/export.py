"""Excel report generation for M08 evaluation runs."""
from __future__ import annotations

from collections import Counter
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.m08_auto_evaluation.services.method_config import (
    METHOD_LABELS,
    METHOD_STANDARD_LABELS,
    TASK_PROFILE_LABELS,
)
from modules.m08_auto_evaluation.services.metrics import aggregate

_EVALUATION_STATUS_LABELS = {
    "scored": "已计算",
    "data_missing": "不适用：缺少数据",
    "unavailable": "不可用：依赖或模型未配置",
    "error": "计算失败",
    "not_selected": "未选择",
}


def _evaluation_status_label(value: str | None) -> str:
    return _EVALUATION_STATUS_LABELS.get(value or "", value or "未计算")


def build_evaluation_workbook(run: dict, results: list[dict]) -> bytes:
    """生成便于测试人员筛选和归因的 Excel 原始报告。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "原始评测报告"
    headers = [
        "序号", "问题", "标准答案", "智能体回答", "得分", "规则评测", "LLM Judge评分", "LLM Judge理由", "状态", "耗时(ms)",
        "维度", "难度", "来源", "归因", "错误信息", "用例ID", "智能体额外输出",
        "标准答案比对状态", "标准答案比对方法", "规则评测状态", "规则评测明细(JSON)", "Judge状态", "Judge判据(JSON)",
        "BLEU", "BLEU状态", "ROUGE-L", "ROUGE-L Precision", "ROUGE-L Recall", "ROUGE-L状态", "改写语义相似度", "改写约束Precision", "改写约束Recall", "改写约束F1",
        "意图标准标签", "意图预测标签", "意图是否正确", "RAG Context Precision", "RAG Context Recall", "RAG Faithfulness", "RAG Answer Relevancy",
        "BLEU计算参数(JSON)", "ROUGE-L计算参数(JSON)", "中间节点输出(JSON)",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="6750A4")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, item in enumerate(results, start=1):
        scores = item.get("scores") or {}
        intermediate_nodes = ((scores.get("intermediate") or {}).get("nodes") or {})
        rewrite = intermediate_nodes.get("rewrite") or {}
        intent = intermediate_nodes.get("intent") or {}
        rag = intermediate_nodes.get("rag") or {}
        rag_metrics = rag.get("metrics") or {}
        bleu = scores.get("bleu") or {}
        rouge_l = scores.get("rouge_l") or {}
        rule = scores.get("rule") or {}
        judge = scores.get("judge") or {}
        answer_comparison = scores.get("answer_comparison") or {}
        sheet.append([
            index,
            item.get("question") or "",
            item.get("gold_answer") or "",
            item.get("answer") or "",
            scores.get("score"),
            rule.get("score"),
            judge.get("score"),
            judge.get("reason"),
            item.get("status") or "",
            scores.get("latency_ms"),
            item.get("dimension") or "",
            item.get("difficulty") or "",
            item.get("source") or "",
            item.get("diagnosis") or "",
            item.get("error_message") or "",
            item.get("case_uid") or "",
            json.dumps(item.get("agent_observations"), ensure_ascii=False) if item.get("agent_observations") else "",
            _evaluation_status_label(answer_comparison.get("status", "scored" if scores.get("score") is not None else "not_selected")),
            answer_comparison.get("method", scores.get("method")),
            _evaluation_status_label(rule.get("status", "scored" if rule.get("score") is not None else "not_selected")),
            json.dumps(rule, ensure_ascii=False) if rule else "",
            _evaluation_status_label(judge.get("status", "scored" if judge.get("score") is not None else "not_selected")),
            json.dumps(judge.get("criteria") or [], ensure_ascii=False) if judge else "",
            bleu.get("score"), _evaluation_status_label(bleu.get("status", "not_selected")),
            rouge_l.get("score"), rouge_l.get("precision"), rouge_l.get("recall"), _evaluation_status_label(rouge_l.get("status", "not_selected")),
            rewrite.get("semantic_similarity"), rewrite.get("constraint_precision"),
            rewrite.get("constraint_recall"), rewrite.get("constraint_f1"),
            intent.get("reference"), intent.get("predicted"), intent.get("correct"),
            rag_metrics.get("context_precision"), rag_metrics.get("context_recall"),
            rag_metrics.get("faithfulness"), rag_metrics.get("answer_relevancy"),
            json.dumps(bleu.get("parameters") or {}, ensure_ascii=False) if bleu else "",
            json.dumps(rouge_l.get("parameters") or {}, ensure_ascii=False) if rouge_l else "",
            json.dumps(intermediate_nodes, ensure_ascii=False) if intermediate_nodes else "",
        ])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [8, 42, 42, 42, 10, 12, 14, 36, 12, 12, 16, 16, 16, 24, 32, 24, 40] + [18] * (len(headers) - 17)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 24

    summary = workbook.create_sheet("运行摘要")
    composition_id = run.get("composition_id")
    task_profile = run.get("task_profile") or "question_answering"
    evaluation_methods = run.get("evaluation_methods")
    if evaluation_methods is None:
        evaluation_methods = ["answer_comparison", "rules"]
        if (run.get("judge_eval") or {}).get("enabled"):
            evaluation_methods.append("llm_as_judge")
    intermediate_config = run.get("intermediate_eval") or {"enabled": False, "nodes": []}
    judge_config = run.get("judge_eval") or {}
    method_labels = [METHOD_LABELS.get(item, item) for item in evaluation_methods]
    node_labels = {"rewrite": "改写", "intent": "意图识别", "rag": "RAG"}
    selected_nodes = intermediate_config.get("nodes") or []
    summary_metrics = aggregate(
        results,
        intermediate_config,
        judge_config,
        task_profile,
        evaluation_methods,
    )
    summary_rows = [
        ("运行名称", run.get("name") or f"运行 #{run.get('run_id')}"),
        ("运行ID", run.get("run_id")),
        ("可执行评测集版本", f"#{composition_id}" if composition_id is not None else "—"),
        ("适配器", run.get("adapter") or ""),
        ("任务类型", TASK_PROFILE_LABELS.get(task_profile, task_profile)),
        ("评测方法", "、".join(method_labels) or "未选择"),
        ("中间过程节点", "、".join(node_labels.get(item, item) for item in selected_nodes) or "未启用"),
        ("Judge附加输入", "、".join(label for enabled, label in [
            (judge_config.get("include_history"), "多轮历史"),
            (judge_config.get("include_intermediate"), "中间节点结果"),
            (judge_config.get("include_retrieved"), "检索结果"),
        ] if enabled) or "无"),
        ("国标标注说明", "国标公式/条文引用不代表本项目通过国标认证；人工 MOS 本版本未纳入。"),
        ("状态", run.get("status") or ""),
        ("结果数量", len(results)),
        ("创建时间", run.get("created_at") or ""),
        ("开始时间", run.get("started_at") or ""),
        ("完成时间", run.get("finished_at") or ""),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 48
    for cell in summary[1]:
        cell.font = Font(bold=True)

    metric_sheet = workbook.create_sheet("指标汇总")
    metric_sheet.append(["节点 / 方法", "指标", "结果", "计算状态", "已评分样本", "适用样本", "国标标注", "计算参数 / 说明"])
    header_fill = PatternFill("solid", fgColor="6750A4")
    for cell in metric_sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    def add_metric(
        group: str,
        metric: str,
        value: object,
        scored: object,
        total: object,
        standard: str,
        params: object = "",
        status: str | None = None,
    ) -> None:
        metric_sheet.append([
            group, metric, value if value is not None else "不适用",
            _evaluation_status_label(status or ("scored" if value is not None else "data_missing")),
            scored, total, standard, params,
        ])
        if isinstance(value, (int, float)) and 0 <= value <= 1:
            metric_sheet.cell(row=metric_sheet.max_row, column=3).number_format = "0.0%"

    scored_answer = [item for item in results if (item.get("scores") or {}).get("score") is not None]
    add_metric(
        "标准答案比对", "平均得分",
        round(sum(float(item["scores"]["score"]) for item in scored_answer) / len(scored_answer), 4) if scored_answer else None,
        len(scored_answer), len(results), "项目方法", "短答案精确匹配；长答案语义相似度",
        "scored" if scored_answer else "data_missing",
    )
    if "rules" in evaluation_methods:
        add_metric("规则评测", "通过率", summary_metrics.get("rule_passed_rate"), summary_metrics.get("rule_scored"), len(results), "项目方法", "must_have_points / acceptable_answers", "scored" if summary_metrics.get("rule_scored") else "data_missing")
    if "llm_as_judge" in evaluation_methods:
        judge_statuses = Counter(
            ((item.get("scores") or {}).get("judge") or {}).get("status", "data_missing")
            for item in results
        )
        judge_status = (
            "scored" if summary_metrics.get("judge_scored") else "error" if judge_statuses.get("error")
            else "unavailable" if judge_statuses.get("unavailable") else "data_missing"
        )
        add_metric("LLM-as-a-Judge", "通过率", summary_metrics.get("judge_passed_rate"), summary_metrics.get("judge_scored"), len(results), "GB/T 45288.2—2025 第 6.5(c) 提及的方法（非指标）", "固定输入问题/标准答案/实际回答；可选上下文见运行摘要", judge_status)
    for method in ("bleu", "rouge_l"):
        if method in evaluation_methods:
            metric = summary_metrics.get("reference_metrics", {}).get(method) or {}
            params = next(
                (
                    ((item.get("scores") or {}).get(method) or {}).get("parameters")
                    for item in results
                    if ((item.get("scores") or {}).get(method) or {}).get("parameters")
                ),
                None,
            )
            add_metric(METHOD_LABELS[method], "平均分", metric.get("score"), metric.get("scored"), metric.get("total"), METHOD_STANDARD_LABELS[method], json.dumps(params or {}, ensure_ascii=False), metric.get("status"))

    intermediate = summary_metrics.get("intermediate") or {}
    for node, metrics in (intermediate.get("nodes") or {}).items():
        node_label = node_labels.get(node, node)
        if node == "rewrite":
            node_metrics = [
                ("语义相似度", "semantic_similarity", "项目方法"),
                ("约束Precision", "constraint_precision", "项目方法"),
                ("约束Recall", "constraint_recall", "项目方法"),
                ("约束F1", "constraint_f1", "项目方法"),
            ]
        elif node == "intent":
            node_metrics = [
                ("Accuracy", "accuracy", "GB/T 45288.2—2025 附录 A.1.1 公式"),
                ("Precision（micro）", "micro_precision", "GB/T 45288.2—2025 附录 A.1.3 公式"),
                ("Recall（micro）", "micro_recall", "GB/T 45288.2—2025 附录 A.1.2 公式"),
                ("Micro-F1", "micro_f1", "GB/T 45288.2—2025 附录 A.1.4 公式"),
                ("Macro-F1", "macro_f1", "项目扩展；非附录 A.1 明示指标"),
            ]
        else:
            node_metrics = [(name, name, "RAGAS 项目方法") for name in (
                "context_precision", "context_recall", "faithfulness", "answer_relevancy"
            )]
        for label, key, standard in node_metrics:
            value = metrics.get(key)
            metric_status = (
                "scored" if value is not None else metrics.get("status")
                if metrics.get("status") in {"error", "unavailable"} else "data_missing"
            )
            add_metric(node_label, label, value, metrics.get("scored"), len(results), standard, status=metric_status)
    metric_sheet.freeze_panes = "A2"
    metric_sheet.auto_filter.ref = metric_sheet.dimensions
    for index, width in enumerate([20, 24, 16, 18, 16, 16, 54, 72], start=1):
        metric_sheet.column_dimensions[get_column_letter(index)].width = width
    for row in metric_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    intent_summary = (intermediate.get("nodes") or {}).get("intent") or {}
    if intent_summary.get("by_label"):
        intent_sheet = workbook.create_sheet("意图分类明细")
        intent_sheet.append(["标准意图标签", "Precision", "Recall", "F1（项目扩展）", "Support", "标准标注"])
        for cell in intent_sheet[1]:
            cell.fill = PatternFill("solid", fgColor="6750A4")
            cell.font = Font(color="FFFFFF", bold=True)
        for label, values in intent_summary["by_label"].items():
            intent_sheet.append([
                label, values.get("precision"), values.get("recall"), values.get("f1"), values.get("support"),
                "Precision / Recall 公式参照 GB/T 45288.2—2025 附录 A.1.3 / A.1.2；类别 F1 为项目扩展。",
            ])
        intent_sheet.append(["混淆矩阵", json.dumps(intent_summary.get("confusion") or {}, ensure_ascii=False), "", "", "", ""])
        intent_sheet.column_dimensions["A"].width = 28
        for column in "BCDE":
            intent_sheet.column_dimensions[column].width = 20
        intent_sheet.column_dimensions["F"].width = 78
        for row in intent_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    for exported_sheet in workbook.worksheets:
        for row in exported_sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
    workbook.save(output)
    return output.getvalue()
